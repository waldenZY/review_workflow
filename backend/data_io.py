import hashlib
import io
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


MISSING = object()


def pointer_escape(value):
    return str(value).replace('~', '~0').replace('/', '~1')


def at(value, path, default=MISSING):
    if path == '':
        return value
    if not path.startswith('/'):
        raise ValueError('字段路径必须使用 JSON Pointer，例如 /answer。')
    try:
        for token in path[1:].split('/'):
            token = token.replace('~1', '/').replace('~0', '~')
            value = value[int(token)] if isinstance(value, list) else value[token]
        return value
    except (KeyError, IndexError, TypeError, ValueError):
        return default


def fields_for(records):
    found = {}

    def visit(value, path):
        if path:
            found.setdefault(path, set()).add(type(value).__name__)
        if isinstance(value, dict):
            for key, child in value.items():
                visit(child, f'{path}/{pointer_escape(key)}')

    for row in records:
        visit(row, '')
    return [{'path': path, 'types': sorted(types)} for path, types in found.items()]


def excel_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def parse_file(filename, content, sheet=None, header_row=1):
    suffix = Path(filename).suffix.lower()
    if suffix not in ('.jsonl', '.json', '.xlsx'):
        raise ValueError('仅支持 JSONL、JSON、XLSX，不支持 CSV。')
    if len(content) > 25 * 1024 * 1024:
        raise ValueError('单文件最大 25 MB。')
    sources = []
    sheet_name = None
    if suffix in ('.json', '.jsonl'):
        text = content.decode('utf-8-sig')
        if suffix == '.json':
            records = json.loads(text)
            if isinstance(records, dict):
                records = [records]
            if not isinstance(records, list):
                raise ValueError('JSON 顶层应是对象或对象数组。')
            sources = list(range(1, len(records) + 1))
        else:
            records = []
            for line_number, line in enumerate(text.splitlines(), 1):
                if line.strip():
                    try:
                        records.append(json.loads(line))
                    except json.JSONDecodeError as exc:
                        raise ValueError(f'JSONL 第 {line_number} 行不是合法 JSON：{exc.msg}') from exc
                    sources.append(line_number)
    else:
        if header_row < 1:
            raise ValueError('表头行必须大于等于 1。')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            sheet_name = sheet or workbook.sheetnames[0]
            if sheet_name not in workbook.sheetnames:
                raise ValueError('工作表不存在。')
            ws = workbook[sheet_name]
            rows = ws.iter_rows(min_row=header_row, values_only=True)
            header = next(rows, None)
            if not header or any(value is None or not str(value).strip() for value in header):
                raise ValueError('表头不能为空；请检查所选表头行及空列。')
            names = [str(value) for value in header]
            if len(names) != len(set(names)):
                raise ValueError('表头不可重复。')
            records = []
            for index, row in enumerate(rows, header_row + 1):
                if all(value is None for value in row):
                    continue
                records.append({name: excel_value(value) for name, value in zip(names, row)})
                sources.append(index)
        finally:
            workbook.close()
    if not records or any(not isinstance(row, dict) for row in records):
        raise ValueError('文件必须包含至少一条对象记录，每条记录均应是对象。')
    # Reject non-standard numeric values before they reach SQLite/API/Excel.
    json.dumps(records, allow_nan=False)
    return records, sources, sheet_name, hashlib.sha256(content).hexdigest()


def flatten(value, prefix=''):
    result = {}
    if isinstance(value, dict) and value:
        for key, child in value.items():
            result.update(flatten(child, f'{prefix}/{pointer_escape(key)}'))
    else:
        result[prefix] = value
    return result


def write_xlsx(path, rows, columns):
    if len(columns) > 16384 or len(rows) > 1048575:
        raise ValueError('超过 XLSX 行列容量，请缩小导出范围。')
    workbook = Workbook()
    ws = workbook.active
    ws.title = '评测结果'
    for row_index, row in enumerate([columns, *rows], 1):
        for col_index, value in enumerate(row, 1):
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            if isinstance(value, str) and len(value) > 32767:
                raise ValueError(f'第 {row_index} 行第 {col_index} 列超过 Excel 单元格限制；完整结果仍保留在运行记录中。')
            cell = ws.cell(row_index, col_index, value)
            if isinstance(value, str):
                cell.data_type = 's'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    workbook.save(path)
    workbook.close()
