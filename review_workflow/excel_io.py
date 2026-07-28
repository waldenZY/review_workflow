from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import RESERVED_COLUMNS


def _header_map(sheet: Any, header_row: int) -> dict[str, int]:
    headers: dict[str, int] = {}
    for cell in sheet[header_row]:
        if cell.value is None:
            continue
        name = str(cell.value)
        if name in headers:
            raise ValueError(f"Duplicate Excel field: {name}")
        headers[name] = cell.column
    return headers


def validate_source(config: dict[str, Any]) -> None:
    source = config["source"]
    workbook = load_workbook(source["file"], read_only=True, data_only=False)
    try:
        if source["sheet"] not in workbook.sheetnames:
            raise ValueError(f"Sheet does not exist: {source['sheet']}")
        headers = _header_map(
            workbook[source["sheet"]], source["header_row"]
        )
        required = [
            item["field"]
            for item in config["fields"]["inputs"]
            + config["fields"]["outputs"]
        ]
        if "key" in config["selection"]:
            required.append(config["selection"]["key"]["field"])
        missing = [name for name in required if name not in headers]
        if missing:
            raise ValueError(f"Missing Excel fields: {', '.join(missing)}")
        collision = RESERVED_COLUMNS.intersection(headers)
        if collision:
            raise ValueError(
                "Source Excel already contains reserved columns: "
                + ", ".join(sorted(collision))
            )
    finally:
        workbook.close()


def open_review_workbook(
    config: dict[str, Any], output_file: Path
) -> tuple[Any, Any, dict[str, int], int, int]:
    source = config["source"]
    output_file.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source["file"], output_file)
    workbook = load_workbook(output_file)
    sheet = workbook[source["sheet"]]
    headers = _header_map(sheet, source["header_row"])
    status_column = sheet.max_column + 1
    output_column = sheet.max_column + 2
    sheet.cell(source["header_row"], status_column, "review_status")
    sheet.cell(source["header_row"], output_column, "review_raw_output")
    return workbook, sheet, headers, status_column, output_column


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def select_cases(
    config: dict[str, Any], sheet: Any, headers: dict[str, int]
) -> list[dict[str, Any]]:
    first_row = config["source"]["header_row"] + 1
    selection = config["selection"]
    if "all" in selection:
        rows = list(range(first_row, sheet.max_row + 1))
    elif "excel_rows" in selection:
        rows = sorted(selection["excel_rows"])
        if any(row > sheet.max_row for row in rows):
            raise ValueError("Selected Excel row is outside the sheet.")
    else:
        key = selection["key"]
        column = headers[key["field"]]
        rows = [
            row
            for row in range(first_row, sheet.max_row + 1)
            if sheet.cell(row, column).value in key["values"]
        ]

    cases = []
    for row in rows:
        configured = (
            config["fields"]["inputs"] + config["fields"]["outputs"]
        )
        if all(
            sheet.cell(row, headers[item["field"]]).value is None
            for item in configured
        ):
            continue

        def fields(items: list[dict[str, str]]) -> list[dict[str, str]]:
            return [
                {
                    "label": item["label"],
                    "value": _cell_text(
                        sheet.cell(row, headers[item["field"]]).value
                    ),
                }
                for item in items
            ]

        cases.append(
            {
                "case_key": f"{sheet.title}#row_{row}",
                "sheet": sheet.title,
                "excel_row": row,
                "inputs": fields(config["fields"]["inputs"]),
                "outputs": fields(config["fields"]["outputs"]),
            }
        )
    if not cases:
        raise ValueError("Selection produced no reviewable rows.")
    return cases
