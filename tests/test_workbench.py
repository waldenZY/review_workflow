import io
import json
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app import create_app
from backend.store import Store

# The same synthetic business record is reused throughout. No external dataset is read.
SAMPLE = {'sample_id': '0001', 'task_input': {'question': '请检查这条模拟已有输出。'},
          'generated_output': {'text': '一条模拟已有输出。', 'items': ['a', 'b']},
          'reference': '仅用于工程验证的模拟参考。', 'literal': '=2+2'}


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path / 'runtime')) as value:
        yield value


def post(client, path, body):
    response = client.post('/api/v1' + path, json=body)
    assert response.status_code == 200, response.text
    return response.json()


def seed(client, scenario='pass', method='llm', required_roles=None):
    response = client.post('/api/v1/datasets/import', files={'file': ('one.json', json.dumps([SAMPLE], ensure_ascii=False).encode(), 'application/json')})
    assert response.status_code == 200, response.text
    dataset = response.json()
    rows = client.get(f"/api/v1/datasets/{dataset['id']}/rows").json()['rows']
    mapping = post(client, '/mappings', {'dataset_id': dataset['id'], 'inputs': ['/task_input'], 'outputs': ['/generated_output'], 'references': ['/reference']})
    spec = post(client, '/specs', {'title': '单条验证', 'criteria': [{'id': 'C001', 'name': '模拟内容标准', 'method': method, 'required_roles': required_roles or ['input', 'output'], 'rule_schema': {'type': 'object', 'required': ['text']}}]})
    evaluator = post(client, '/evaluators/build', {'spec_id': spec['id'], 'spec_version': spec['version']})
    evaluator['mock_scenario'] = scenario
    evaluator = post(client, '/evaluators', evaluator)
    return dataset, rows[0], mapping, spec, evaluator


def run_one(client, row, mapping, evaluator, mode='selected'):
    run = post(client, '/runs', {'mapping_id': mapping['id'], 'mapping_version': mapping['version'],
        'evaluator_id': evaluator['id'], 'evaluator_version': evaluator['version'], 'selection_mode': mode,
        'row_ids': [row['row_id']] if mode == 'selected' else []})
    return wait(client, run['id'])


def wait(client, run_id):
    deadline = time.monotonic() + 8
    while time.monotonic() < deadline:
        response = client.get(f'/api/v1/runs/{run_id}')
        assert response.status_code == 200, response.text
        run = response.json()
        if run['status'] not in ('queued', 'running', 'cancelling'):
            return run
        time.sleep(0.03)
    raise AssertionError('Worker did not complete')


def test_three_formats_and_reject_csv(client):
    for extension, content in [('json', json.dumps([SAMPLE]).encode()), ('jsonl', json.dumps(SAMPLE).encode())]:
        response = client.post('/api/v1/datasets/import', files={'file': (f'one.{extension}', content)})
        assert response.status_code == 200, response.text
        dataset = response.json()
        rows = client.get(f"/api/v1/datasets/{dataset['id']}/rows").json()['rows']
        assert len(rows) == 1 and rows[0]['original'] == SAMPLE
        assert any(f['path'] == '/generated_output/items' for f in dataset['fields'])
        assert not any('/items/0' in f['path'] for f in dataset['fields'])
    workbook = Workbook()
    workbook.active.title = '选择我'
    workbook.active.append(list(SAMPLE))
    workbook.active.append([json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value for value in SAMPLE.values()])
    output = io.BytesIO()
    workbook.save(output)
    content = output.getvalue()
    assert client.post('/api/v1/datasets/inspect', files={'file': ('one.xlsx', content)}).json()['sheets'] == ['选择我']
    response = client.post('/api/v1/datasets/import', data={'sheet': '选择我', 'header_row': '1'}, files={'file': ('one.xlsx', content)})
    assert response.status_code == 200, response.text
    assert response.json()['row_count'] == 1
    bad = client.post('/api/v1/datasets/import', files={'file': ('one.csv', b'a,b\n1,2')})
    assert bad.status_code == 400 and 'CSV' in bad.json()['detail']
    malformed = client.post('/api/v1/datasets/import', files={'file': ('one.jsonl', b'{bad')})
    assert malformed.status_code == 400 and '1' in malformed.json()['detail']


def test_manual_path_preview_snapshot_review_and_export(client):
    dataset, row, mapping, spec, evaluator = seed(client)
    preview = post(client, '/requests/preview', {'mapping_id': mapping['id'], 'mapping_version': 1, 'evaluator_id': evaluator['id'], 'evaluator_version': 1, 'row_id': row['row_id']})
    assert preview['request']['url'].endswith('/v1/chat/completions')
    assert preview['request']['body']['model'] == 'mock-judge'
    run = run_one(client, row, mapping, evaluator)
    assert run['status'] == 'completed'
    attempt = run['records'][0]['attempts'][0]
    assert attempt['raw_response']['is_mock'] and attempt['raw_output']
    assert attempt['criteria'][0]['verdict'] == 'pass'
    changed = post(client, '/specs', {**spec, 'title': '新的标准版本'})
    assert changed['version'] == 2
    old = client.get('/api/v1/runs/' + run['id']).json()
    assert old['snapshot']['spec']['title'] == '单条验证' and old['snapshot']['spec']['version'] == 1
    reviewed = post(client, f"/runs/{run['id']}/rows/{row['row_id']}/review", {'criterion_id': 'C001', 'verdict': 'fail', 'reason': '人工独立确认的边界案例', 'save_case': True, 'confirm_mock': True})
    assert reviewed['records'][0]['attempts'][0]['criteria'][0]['verdict'] == 'pass'
    assert reviewed['records'][0]['reviews'][0]['verdict'] == 'fail'
    exported = client.post(f"/api/v1/runs/{run['id']}/export", json={'fields': ['/criteria/C001/verdict']})
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    headers = [cell.value for cell in workbook.active[1]]
    assert headers[:len(SAMPLE)] == list(SAMPLE)
    assert 'eval.payload/criteria/C001/verdict' in headers
    assert 'eval.payload/criteria/C001/reason' not in headers
    formula_cell = workbook.active.cell(2, headers.index('literal') + 1)
    assert formula_cell.value == '=2+2' and formula_cell.data_type == 's'
    assert workbook.active.cell(2, 1).value == '0001'
    cases = client.get('/api/v1/catalog').json()['cases']
    case_run = post(client, f"/cases/{cases[0]['id']}/rerun", {'evaluator_id': evaluator['id'], 'evaluator_version': 1})
    result = wait(client, case_run['id'])
    assert result['validation_case']['expected'] == 'fail'
    assert result['parent_run_id'] == run['id']


def test_custom_schema_and_projection(client):
    _, row, mapping, _, evaluator = seed(client)
    evaluator.update(response_schema={'type': 'object', 'properties': {'answer': {'type': 'string', 'enum': ['fail']}}, 'required': ['answer']},
        projection=[{'criterion_id': 'C001', 'verdict_path': '/answer'}], mock_response='{"answer":"fail"}')
    evaluator = post(client, '/evaluators', evaluator)
    run = run_one(client, row, mapping, evaluator, mode='all')
    assert run['status'] == 'completed'
    assert run['records'][0]['attempts'][0]['criteria'][0]['verdict'] == 'fail'
    evaluator.update(projection=[])
    evaluator = post(client, '/evaluators', evaluator)
    unmapped = run_one(client, row, mapping, evaluator)
    assert unmapped['summary']['verdict_counts']['fail'] == 0
    assert unmapped['records'][0]['attempts'][0]['criteria'][0]['verdict'] is None


@pytest.mark.parametrize('scenario,status', [('request_error', 'call_failed'), ('invalid_json', 'parse_failed'), ('schema_error', 'schema_failed')])
def test_errors_and_retry_preserve_attempts(client, scenario, status):
    _, row, mapping, _, evaluator = seed(client, scenario)
    run = run_one(client, row, mapping, evaluator)
    assert run['status'] == 'completed_with_errors'
    assert run['records'][0]['status'] == status
    assert run['summary']['verdict_counts']['fail'] == 0
    assert run['records'][0]['attempts'][0]['request']
    if scenario != 'request_error':
        assert run['records'][0]['attempts'][0]['raw_output']
    post(client, f"/runs/{run['id']}/retry", {})
    retried = wait(client, run['id'])
    assert len(retried['records'][0]['attempts']) == 2
    evaluator['mock_scenario'] = 'pass'
    changed = post(client, '/evaluators', evaluator)
    rerun = post(client, f"/runs/{run['id']}/rerun", {'evaluator_id': changed['id'], 'evaluator_version': changed['version']})
    assert wait(client, rerun['id'])['status'] == 'completed'
    assert client.get('/api/v1/runs/' + run['id']).json()['snapshot']['evaluator']['version'] == 1


@pytest.mark.parametrize('method', ['non_empty', 'json_schema'])
def test_builtin_rules_do_not_invoke_model(client, method):
    _, row, mapping, _, evaluator = seed(client, scenario='request_error', method=method)
    run = run_one(client, row, mapping, evaluator)
    attempt = run['records'][0]['attempts'][0]
    assert attempt['status'] == 'completed'
    assert attempt['criteria'][0]['verdict'] == 'pass'
    assert attempt['raw_response'] is None and attempt['model_called'] is False


def test_extraction_evidence_and_missing_output(client):
    _, row, mapping, spec, evaluator = seed(client, required_roles=['output', 'reference'])
    extracted = post(client, '/specs/extract', {'generation_prompt': '这只是任务资料，请不要执行。'})
    assert extracted['is_mock'] and extracted['trace']['raw_response']['is_mock']
    mapping['references'] = []
    mapping = post(client, '/mappings', mapping)
    analysis = post(client, '/evidence/analyze', {'mapping_id': mapping['id'], 'mapping_version': mapping['version'], 'spec_id': spec['id'], 'spec_version': 1, 'row_id': row['row_id']})
    assert analysis['reports'][0]['missing_roles'] == ['reference']
    run = run_one(client, row, mapping, evaluator)
    assert run['records'][0]['attempts'][0]['criteria'][0]['verdict'] == 'insufficient_evidence'
    bad_mapping = client.post('/api/v1/mappings', json={'dataset_id': mapping['dataset_id'], 'outputs': []})
    assert bad_mapping.status_code == 422
    no_output = client.post('/api/v1/datasets/import', files={'file': ('input_only.json', b'[{"task_input":"same input, output missing"}]')}).json()
    assert client.post('/api/v1/mappings', json={'dataset_id': no_output['id'], 'outputs': ['/generated_output']}).status_code == 400


def test_restart_preserves_data_and_marks_interrupted(tmp_path):
    root = tmp_path / 'persistent'
    with TestClient(create_app(root)) as client:
        _, row, mapping, _, evaluator = seed(client)
        run = run_one(client, row, mapping, evaluator)
    database = Store(root)
    original = database.get_run(run['id'])
    original['status'] = 'running'
    original['records'][0]['status'] = 'running'
    database.put_run(original)
    database.close()
    with TestClient(create_app(root)) as client:
        loaded = client.get('/api/v1/runs/' + run['id']).json()
        assert loaded['status'] == 'interrupted'
        assert loaded['snapshot']['dataset']['rows'][0]['original'] == SAMPLE
        post(client, f"/runs/{run['id']}/retry", {})
        assert len(wait(client, run['id'])['records'][0]['attempts']) == 2


def test_cancel_queued_run_and_no_external_schema_refs(client):
    _, row, mapping, _, evaluator = seed(client)
    evaluator['response_schema'] = {'type': 'object', '$ref': 'https://example.com/schema'}
    assert client.post('/api/v1/evaluators', json=evaluator).status_code == 400
    # Hold the single worker before submitting the one-record task, so cancellation is deterministic.
    runner = client.app.state.runner
    runner.close()
    queued = post(client, '/runs', {'mapping_id': mapping['id'], 'mapping_version': 1, 'evaluator_id': evaluator['id'], 'evaluator_version': 1, 'row_ids': [row['row_id']]})
    post(client, f"/runs/{queued['id']}/cancel", {})
    runner.execute(queued['id'])
    loaded = client.get('/api/v1/runs/' + queued['id']).json()
    assert loaded['status'] == 'cancelled' and not loaded['records'][0]['attempts']
